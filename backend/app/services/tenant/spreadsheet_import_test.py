"""Tests for reading a spreadsheet file into the workbook shape.

The ones that matter most go out through the renderer and back in, because
import and export only stay agreed about how a fill or a currency is spelled
while something checks that they do.
"""

import pytest

from app.core.messages import FileMessages
from app.services.export.spreadsheet import render_xlsx
from app.services.tenant.files_spreadsheet import (
    FileContentError,
    normalize_spreadsheet_content,
)
from app.services.tenant.spreadsheet_import import decode_text, parse_spreadsheet_file


def workbook(**sheet) -> dict:
    """A canonical one-sheet workbook, normalized the way a stored one is."""
    base = {
        "id": "s1",
        "name": "Budget",
        "dimensions": {"rows": 5, "cols": 3},
        "cells": {},
        "columns": {},
        "rows": {},
        "cellStyles": {},
        "frozen": {"rows": 0, "cols": 0},
    }
    base.update(sheet)
    return normalize_spreadsheet_content(
        {"schema_version": 3, "kind": "spreadsheet", "sheets": [base]}
    )


# ── CSV ──────────────────────────────────────────────────────────────────────


def test_a_csv_becomes_one_sheet_named_after_its_file() -> None:
    sheets = parse_spreadsheet_file("Q1 sales.csv", b"a,b\n1,2\n")

    assert len(sheets) == 1
    assert sheets[0]["name"] == "Q1 sales"


def test_numbers_arrive_as_numbers_and_text_as_text() -> None:
    sheets = parse_spreadsheet_file("x.csv", b"Widget,3,4.5,-2\n")

    assert sheets[0]["cells"] == {
        "0:0": "Widget",
        "0:1": 3,
        "0:2": 4.5,
        "0:3": -2,
    }


def test_a_formula_stays_a_formula() -> None:
    """``=`` text is how this app stores a formula; the grid evaluates it."""
    sheets = parse_spreadsheet_file("x.csv", b"=1+1\n")

    assert sheets[0]["cells"]["0:0"] == "=1+1"


def test_empty_fields_leave_no_cell() -> None:
    sheets = parse_spreadsheet_file("x.csv", b"a,,c\n")

    assert sheets[0]["cells"] == {"0:0": "a", "0:2": "c"}


def test_a_tsv_is_split_on_tabs() -> None:
    sheets = parse_spreadsheet_file("x.tsv", b"a\tb\n")

    assert sheets[0]["cells"] == {"0:0": "a", "0:1": "b"}


def test_a_file_that_is_not_a_spreadsheet_is_refused() -> None:
    with pytest.raises(FileContentError) as excinfo:
        parse_spreadsheet_file("notes.pdf", b"%PDF-1.4")

    assert excinfo.value.code == FileMessages.SPREADSHEET_UNREADABLE_FILE


def test_an_xlsx_that_is_not_one_is_refused() -> None:
    with pytest.raises(FileContentError):
        parse_spreadsheet_file("broken.xlsx", b"not a zip at all")


def test_text_that_is_not_utf8_still_imports() -> None:
    """A file from an older tool is worth a mangled accent, not a refusal."""
    sheets = parse_spreadsheet_file("x.csv", "café".encode("cp1252"))

    assert sheets[0]["cells"]["0:0"].startswith("caf")


# ── XLSX, through the renderer and back ──────────────────────────────────────


def test_values_survive_the_round_trip() -> None:
    content = workbook(
        cells={"0:0": "Item", "1:0": "Widget", "1:1": 12.5, "2:1": "=B2*2"}
    )

    back = parse_spreadsheet_file("b.xlsx", render_xlsx(content, title="Budget"))

    assert back[0]["cells"] == content["sheets"][0]["cells"]


def test_a_cells_look_survives_the_round_trip() -> None:
    content = workbook(
        cells={"0:0": "Item"},
        cellStyles={
            "0:0": {
                "style": {
                    "bold": True,
                    "italic": True,
                    "fill": "#ffee88",
                    "color": "#223344",
                    "align": "center",
                    "valign": "middle",
                }
            }
        },
    )

    back = parse_spreadsheet_file("b.xlsx", render_xlsx(content, title="Budget"))

    assert back[0]["cellStyles"]["0:0"]["style"] == {
        "bold": True,
        "italic": True,
        "fill": "#ffee88",
        "color": "#223344",
        "align": "center",
        "valign": "middle",
    }


def test_a_number_format_survives_the_round_trip() -> None:
    content = workbook(
        cells={"0:0": 12.5},
        cellStyles={
            "0:0": {"format": {"type": "currency", "currency": "USD", "decimals": 2}}
        },
    )

    back = parse_spreadsheet_file("b.xlsx", render_xlsx(content, title="Budget"))
    fmt = back[0]["cellStyles"]["0:0"]["format"]

    assert fmt["type"] == "currency"
    assert fmt["currency"] == "USD"
    assert fmt["decimals"] == 2


def test_a_percent_keeps_its_decimals() -> None:
    content = workbook(
        cells={"0:0": 0.25},
        cellStyles={"0:0": {"format": {"type": "percent", "decimals": 1}}},
    )

    back = parse_spreadsheet_file("b.xlsx", render_xlsx(content, title="Budget"))

    assert back[0]["cellStyles"]["0:0"]["format"] == {
        "type": "percent",
        "decimals": 1,
    }


def test_frozen_panes_and_column_widths_survive() -> None:
    content = workbook(
        cells={"0:0": "Item"},
        columns={"0": {"width": 140}},
        frozen={"rows": 1, "cols": 2},
    )

    back = parse_spreadsheet_file("b.xlsx", render_xlsx(content, title="Budget"))

    assert back[0]["frozen"] == {"rows": 1, "cols": 2}
    assert back[0]["columns"]["0"]["width"] == 140


def test_an_imported_sheet_has_room_to_work_in() -> None:
    """A three-row CSV should not open as a three-row grid."""
    sheets = parse_spreadsheet_file("x.csv", b"a\nb\nc\n")

    assert sheets[0]["dimensions"]["rows"] >= 100
    assert sheets[0]["dimensions"]["cols"] >= 26


def test_a_cell_nobody_styled_carries_no_style() -> None:
    """Excel writes a font on every cell; carrying that over would pin a whole
    imported sheet to one size and colour."""
    content = workbook(cells={"0:0": "plain", "1:1": 3})

    back = parse_spreadsheet_file("b.xlsx", render_xlsx(content, title="Budget"))

    assert back[0]["cellStyles"] == {}


def test_every_tab_of_a_workbook_arrives() -> None:
    content = normalize_spreadsheet_content(
        {
            "schema_version": 3,
            "kind": "spreadsheet",
            "sheets": [
                {"id": "a", "name": "Q1", "cells": {"0:0": "one"}},
                {"id": "b", "name": "Q2", "cells": {"0:0": "two"}},
            ],
        }
    )

    back = parse_spreadsheet_file("y.xlsx", render_xlsx(content, title="Year"))

    assert [s["name"] for s in back] == ["Q1", "Q2"]
    assert [s["cells"]["0:0"] for s in back] == ["one", "two"]


# ── what a field means ───────────────────────────────────────────────────────


def test_a_leading_zero_keeps_its_field_as_text() -> None:
    """``00123`` is a part number or a postcode far more often than it is the
    number 123, and making it one cannot be undone."""
    sheets = parse_spreadsheet_file("x.csv", b"00123,0,0.5,123\n")

    assert sheets[0]["cells"] == {
        "0:0": "00123",
        "0:1": 0,
        "0:2": 0.5,
        "0:3": 123,
    }


def test_true_and_false_arrive_as_booleans() -> None:
    sheets = parse_spreadsheet_file("x.csv", b"TRUE,false,Maybe\n")

    assert sheets[0]["cells"] == {"0:0": True, "0:1": False, "0:2": "Maybe"}


def test_a_field_reads_the_same_from_a_file_as_from_the_clipboard() -> None:
    """The rule lives twice — here and in ``coerceScalar`` on the client — so
    this is the check that the two still say the same thing."""
    sheets = parse_spreadsheet_file("x.csv", b" 42 ,1e3,-7,=A1,.7,-.25,1e309\n")

    assert sheets[0]["cells"] == {
        "0:0": 42,
        "0:1": 1000.0,
        "0:2": -7,
        "0:3": "=A1",
        "0:4": 0.7,
        "0:5": -0.25,
        "0:6": "1e309",
    }


# ── refusing what cannot be carried whole ────────────────────────────────────


def test_a_csv_with_too_many_cells_is_refused_rather_than_trimmed(monkeypatch) -> None:
    """A workbook that came back missing everything past some line, reported
    as imported, is worse than one that did not come back."""
    from app.services.tenant import spreadsheet_import

    # A budget small enough to trip on a file this test can afford to build.
    monkeypatch.setattr(spreadsheet_import, "MAX_IMPORT_CELLS", 500)
    row = b",".join(b"x" for _ in range(50)) + b"\n"

    with pytest.raises(FileContentError) as excinfo:
        parse_spreadsheet_file("big.csv", row * 40)

    assert excinfo.value.code == FileMessages.SPREADSHEET_FILE_TOO_LARGE


def test_a_workbook_with_too_many_tabs_is_refused() -> None:
    from openpyxl import Workbook

    import io as _io

    book = Workbook()
    for index in range(70):
        book.create_sheet(f"S{index}")
    buffer = _io.BytesIO()
    book.save(buffer)

    with pytest.raises(FileContentError) as excinfo:
        parse_spreadsheet_file("many.xlsx", buffer.getvalue())

    assert excinfo.value.code == FileMessages.SPREADSHEET_FILE_TOO_LARGE


def test_a_workbook_declaring_more_than_one_import_reads_is_refused(
    monkeypatch,
) -> None:
    """Judged from the sizes the package declares, before the workbook opens."""
    import io as _io

    from openpyxl import Workbook

    from app.services.tenant import spreadsheet_import

    buffer = _io.BytesIO()
    Workbook().save(buffer)
    monkeypatch.setattr(spreadsheet_import, "MAX_IMPORT_XLSX_BYTES", 1024)
    monkeypatch.setattr(
        spreadsheet_import,
        "load_workbook",
        lambda *a, **k: pytest.fail("the workbook was opened"),
    )

    with pytest.raises(FileContentError) as excinfo:
        parse_spreadsheet_file("big.xlsx", buffer.getvalue())

    assert excinfo.value.code == FileMessages.SPREADSHEET_FILE_TOO_LARGE


def test_a_sheet_claiming_more_rows_than_the_grid_is_refused() -> None:
    """The shape is judged from what the sheet declares, before anything in it
    is read — a file that would take too long never starts."""
    import io as _io

    from openpyxl import Workbook

    book = Workbook()
    book.active["A200000"] = 1  # one cell, a rectangle past the grid
    buffer = _io.BytesIO()
    book.save(buffer)

    with pytest.raises(FileContentError) as excinfo:
        parse_spreadsheet_file("sparse.xlsx", buffer.getvalue())

    assert excinfo.value.code == FileMessages.SPREADSHEET_FILE_TOO_LARGE


def test_a_sparse_sheet_with_a_distant_corner_is_refused() -> None:
    """Two cells can declare a rectangle of millions of coordinates. The cost
    is in the rectangle, not in what is stored."""
    import io as _io

    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet["A1"] = 1
    sheet.cell(row=50_000, column=100, value=1)  # 5,000,000 coordinates
    buffer = _io.BytesIO()
    book.save(buffer)

    with pytest.raises(FileContentError) as excinfo:
        parse_spreadsheet_file("sparse.xlsx", buffer.getvalue())

    assert excinfo.value.code == FileMessages.SPREADSHEET_FILE_TOO_LARGE


def _ods(rows: list[list[str]], *, name: str = "Budget") -> bytes:
    """An OpenDocument spreadsheet holding ``rows`` of text and numbers."""
    import io
    import zipfile

    def cell(value: str) -> str:
        if value.startswith("of:="):
            return (
                f'<table:table-cell table:formula="{value}" office:value-type="float" '
                'office:value="0"><text:p>0</text:p></table:table-cell>'
            )
        if value == "":
            return '<table:table-cell table:number-columns-repeated="2"/>'
        if value.replace(".", "", 1).isdigit():
            return (
                f'<table:table-cell office:value-type="float" '
                f'office:value="{value}"><text:p>{value}</text:p></table:table-cell>'
            )
        return (
            '<table:table-cell office:value-type="string">'
            f"<text:p>{value}</text:p></table:table-cell>"
        )

    body = "".join(
        "<table:table-row>" + "".join(cell(v) for v in row) + "</table:table-row>"
        for row in rows
    )
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">'
        f'<office:body><office:spreadsheet><table:table table:name="{name}">'
        f"{body}</table:table></office:spreadsheet></office:body>"
        "</office:document-content>"
    )
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">'
        '<manifest:file-entry manifest:full-path="/" '
        'manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>'
        '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
        "</manifest:manifest>"
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as package:
        package.writestr(
            "mimetype",
            "application/vnd.oasis.opendocument.spreadsheet",
            compress_type=zipfile.ZIP_STORED,
        )
        package.writestr("content.xml", content)
        package.writestr("META-INF/manifest.xml", manifest)
    return out.getvalue()


def test_an_ods_is_read_in_a_worker_as_its_values() -> None:
    (sheet,) = parse_spreadsheet_file(
        "budget.ods", _ods([["Item", "Cost"], ["Tea", "3"]])
    )

    assert sheet["name"] == "Budget"
    assert sheet["cells"] == {"0:0": "Item", "0:1": "Cost", "1:0": "Tea", "1:1": 3}


def test_an_ods_keeps_its_formulas_where_they_were() -> None:
    # Two empty cells, written as one repeated, put the formula in column C.
    rows = [["2", "3"], ["", "of:=SUM([.A1:.B1])"]]

    (sheet,) = parse_spreadsheet_file("budget.ods", _ods(rows))

    assert sheet["cells"]["1:2"] == "=SUM(A1:B1)"


def test_a_file_calamine_cannot_read_is_refused() -> None:
    with pytest.raises(FileContentError) as refused:
        parse_spreadsheet_file("broken.xls", b"not a workbook")

    assert refused.value.code == FileMessages.SPREADSHEET_UNREADABLE_FILE


@pytest.mark.parametrize(
    ("data", "text"),
    [
        ("café".encode(), "café"),
        ("café".encode("cp1252"), "café"),
        ("café".encode("utf-16"), "café"),
        (b"caf\x81", None),
        (b"a\x00b", None),
    ],
    ids=["utf-8", "cp1252", "utf-16", "no-encoding", "binary"],
)
def test_text_is_read_whole_or_not_at_all(data, text) -> None:
    assert decode_text(data) == text


def test_a_strict_read_refuses_a_table_it_cannot_read_whole() -> None:
    with pytest.raises(FileContentError) as refused:
        parse_spreadsheet_file("broken.csv", b"a,b\nc\x00,d\n", strict=True)

    assert refused.value.code == FileMessages.SPREADSHEET_UNREADABLE_FILE
