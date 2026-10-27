"""Reading Word and OpenDocument text files into the walker's XHTML."""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

from app.core.file_types import FileType
from app.core.messages import FileMessages
from app.services.tenant import document_import
from app.services.tenant.document_import import read_document
from app.services.tenant.files import FileContentError


def _png() -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (40, 20), "red").save(out, "PNG")
    return out.getvalue()


def _saved(document) -> bytes:
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _docx(body: str) -> bytes:
    """A package holding only ``body`` as its document's body."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as package:
        package.writestr(
            "word/document.xml",
            f"<w:document {_W}><w:body>{body}</w:body></w:document>",
        )
    return out.getvalue()


_ODF = " ".join(
    f'xmlns:{prefix}="urn:oasis:names:tc:opendocument:xmlns:{name}"'
    for prefix, name in (
        ("office", "office:1.0"),
        ("style", "style:1.0"),
        ("text", "text:1.0"),
        ("table", "table:1.0"),
        ("draw", "drawing:1.0"),
        ("fo", "xsl-fo-compatible:1.0"),
        ("svg", "svg-compatible:1.0"),
    )
)


def _odt(text: str, *, automatic: str = "", pictures: dict[str, bytes] = {}) -> bytes:
    """An ``.odt`` as an office suite writes one: its mimetype first and stored,
    common styles in ``styles.xml``, automatic ones in ``content.xml``."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as package:
        package.writestr(
            "mimetype",
            "application/vnd.oasis.opendocument.text",
            compress_type=zipfile.ZIP_STORED,
        )
        package.writestr(
            "content.xml",
            f'<office:document-content {_ODF} xmlns:xlink="http://www.w3.org/1999/xlink">'
            f"<office:automatic-styles>{automatic}</office:automatic-styles>"
            f"<office:body><office:text>{text}</office:text></office:body>"
            "</office:document-content>",
        )
        package.writestr(
            "styles.xml",
            f"<office:document-styles {_ODF}><office:styles>"
            '<style:style style:name="Heading" style:family="paragraph"/>'
            '<style:style style:name="Title" style:family="paragraph"'
            ' style:parent-style-name="Heading">'
            '<style:text-properties fo:font-weight="bold"/></style:style>'
            "</office:styles></office:document-styles>",
        )
        for name, data in pictures.items():
            package.writestr(name, data)
    return out.getvalue()


def _body(read) -> str:
    return read.xhtml.removeprefix("<div>").removesuffix("</div>")


# ── Word ─────────────────────────────────────────────────────────────────────


def test_a_word_file_reads_as_headings_formatting_lists_and_links() -> None:
    document = Document()
    document.add_heading("Fete plan", 0)
    document.add_heading("Stalls", 2)
    line = document.add_paragraph("Bring ")
    line.add_run("cakes").bold = True
    line.add_run(" and ")
    jam = line.add_run("jam")
    jam.italic = jam.underline = True
    document.add_paragraph("Tombola", style="List Bullet")
    document.add_paragraph("Raffle tickets", style="List Bullet 2")
    document.add_paragraph("Bins", style="List Bullet")
    document.add_paragraph("First", style="List Number")
    document.add_paragraph("Second", style="List Number")
    linked = document.add_paragraph("See ")
    link = OxmlElement("w:hyperlink")
    link.set(
        qn("r:id"),
        linked.part.relate_to(
            "https://example.org/fete", RELATIONSHIP_TYPE.HYPERLINK, is_external=True
        ),
    )
    run = OxmlElement("w:r")
    words = OxmlElement("w:t")
    words.text = "the site"
    run.append(words)
    link.append(run)
    linked._p.append(link)

    read = read_document(FileType.docx, _saved(document))

    assert _body(read) == (
        "<h1>Fete plan</h1><h2>Stalls</h2>"
        "<p>Bring <strong>cakes</strong> and <em><u>jam</u></em></p>"
        "<ul><li><p>Tombola</p><ul><li><p>Raffle tickets</p></li></ul></li>"
        "<li><p>Bins</p></li></ul>"
        "<ol><li><p>First</p></li><li><p>Second</p></li></ol>"
        '<p>See <a href="https://example.org/fete">the site</a></p>'
    )
    assert read.pictures == {}


def test_a_paragraph_style_formats_its_text_and_a_run_can_undo_it() -> None:
    document = Document()
    shout = document.styles.add_style("Shout", WD_STYLE_TYPE.PARAGRAPH)
    shout.font.bold = True
    line = document.add_paragraph("Loud ", style="Shout")
    line.add_run("quiet").bold = False

    read = read_document(FileType.docx, _saved(document))

    assert _body(read) == "<p><strong>Loud </strong>quiet</p>"


def _numbered(num_id: str, text: str) -> str:
    return (
        f'<w:p><w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/>'
        f"</w:numPr></w:pPr><w:r><w:t>{text}</w:t></w:r></w:p>"
    )


def test_separate_word_lists_stay_separate() -> None:
    read = read_document(
        FileType.docx,
        _docx(_numbered("1", "a") + _numbered("1", "b") + _numbered("2", "c")),
    )

    assert _body(read) == (
        "<ul><li><p>a</p></li><li><p>b</p></li></ul><ul><li><p>c</p></li></ul>"
    )


def test_a_word_table_keeps_its_merged_cells_and_header_row() -> None:
    document = Document()
    table = document.add_table(rows=3, cols=3)
    for row in range(3):
        for column in range(3):
            table.cell(row, column).text = f"{row}{column}"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 2).merge(table.cell(2, 2))

    read = read_document(FileType.docx, _saved(document))

    assert _body(read) == (
        "<table>"
        '<tr><th colspan="2"><p>00</p><p>01</p></th><th><p>02</p></th></tr>'
        '<tr><td><p>10</p></td><td><p>11</p></td><td rowspan="2"><p>12</p><p>22</p></td></tr>'
        "<tr><td><p>20</p></td><td><p>21</p></td></tr>"
        "</table>"
    )


def test_a_word_picture_is_named_by_its_part_and_handed_back() -> None:
    document = Document()
    document.add_picture(io.BytesIO(_png()), width=Inches(1))

    read = read_document(FileType.docx, _saved(document))

    assert _body(read) == (
        '<p><ac:image ac:width="96" ac:height="48">'
        '<ri:attachment ri:filename="word/media/image1.png" /></ac:image></p>'
    )
    assert read.pictures == {"word/media/image1.png": _png()}


def test_tracked_changes_read_accepted_and_a_contents_table_is_left_out() -> None:
    read = read_document(
        FileType.docx,
        _docx(
            "<w:sdt><w:sdtPr><w:docPartObj><w:docPartGallery"
            ' w:val="Table of Contents"/></w:docPartObj></w:sdtPr>'
            "<w:sdtContent><w:p><w:r><w:t>Contents</w:t></w:r></w:p>"
            "</w:sdtContent></w:sdt>"
            "<w:p><w:r><w:t xml:space='preserve'>Tea </w:t></w:r>"
            "<w:ins><w:r><w:t>at four</w:t></w:r></w:ins>"
            "<w:del><w:r><w:delText>at three</w:delText></w:r></w:del></w:p>"
            '<w:p><w:pPr><w:outlineLvl w:val="0"/></w:pPr>'
            "<w:r><w:t>Outlined</w:t></w:r></w:p>"
        ),
    )

    assert _body(read) == "<p>Tea at four</p><h1>Outlined</h1>"


def test_a_table_nested_past_the_limit_is_kept_as_its_rows_text() -> None:
    cell = "<w:tc><w:p><w:r><w:t>deep</w:t></w:r></w:p></w:tc>"
    for _ in range(document_import._MAX_TABLE_DEPTH):
        cell = f"<w:tc><w:tbl><w:tr>{cell}</w:tr></w:tbl></w:tc>"

    read = read_document(FileType.docx, _docx(f"<w:tbl><w:tr>{cell}</w:tr></w:tbl>"))

    assert read.xhtml.count("<table>") == document_import._MAX_TABLE_DEPTH
    assert "<p>deep</p>" in read.xhtml


# ── OpenDocument ─────────────────────────────────────────────────────────────


def test_an_opendocument_file_reads_as_its_styles_lists_and_tables_say() -> None:
    read = read_document(
        FileType.odt,
        _odt(
            '<text:p text:style-name="Title">Fete plan</text:p>'
            '<text:h text:outline-level="2">Stalls</text:h>'
            '<text:p text:style-name="P1">All bold</text:p>'
            '<text:p>Bring <text:span text:style-name="T1">jam</text:span>'
            ' x<text:span text:style-name="T2">2</text:span>'
            "<text:note><text:note-body><text:p>A note</text:p></text:note-body>"
            '</text:note> and <text:a xlink:href="https://example.org/">a link'
            "</text:a></text:p>"
            '<text:list text:style-name="L1"><text:list-item><text:p>One</text:p>'
            "<text:list><text:list-item><text:p>One a</text:p></text:list-item>"
            "</text:list></text:list-item>"
            "<text:list-item><text:p>Two</text:p></text:list-item></text:list>"
            "<table:table><table:table-header-rows><table:table-row>"
            "<table:table-cell><text:p>H1</text:p></table:table-cell>"
            "<table:table-cell><text:p>H2</text:p></table:table-cell>"
            "</table:table-row></table:table-header-rows><table:table-row>"
            '<table:table-cell table:number-columns-spanned="2"><text:p>wide'
            "</text:p></table:table-cell><table:covered-table-cell/>"
            "</table:table-row></table:table>"
            "<text:table-of-content><text:index-body><text:p>Contents</text:p>"
            "</text:index-body></text:table-of-content>",
            automatic=(
                '<style:style style:name="P1" style:family="paragraph">'
                '<style:text-properties fo:font-weight="bold"/></style:style>'
                '<style:style style:name="T1" style:family="text">'
                '<style:text-properties fo:font-style="italic"/></style:style>'
                '<style:style style:name="T2" style:family="text">'
                '<style:text-properties style:text-position="super 58%"/>'
                "</style:style>"
                '<text:list-style style:name="L1">'
                '<text:list-level-style-number text:level="1" style:num-format="1"/>'
                '<text:list-level-style-bullet text:level="2"/></text:list-style>'
            ),
        ),
    )

    assert _body(read) == (
        "<h1>Fete plan</h1><h2>Stalls</h2><p><strong>All bold</strong></p>"
        "<p>Bring <em>jam</em> x<sup>2</sup> and "
        '<a href="https://example.org/">a link</a></p>'
        "<ol><li><p>One</p><ul><li><p>One a</p></li></ul></li>"
        "<li><p>Two</p></li></ol>"
        "<table><tr><th><p>H1</p></th><th><p>H2</p></th></tr>"
        '<tr><td colspan="2"><p>wide</p></td></tr></table>'
    )


def test_separate_opendocument_lists_stay_separate_unless_one_continues() -> None:
    item = "<text:list-item><text:p>{}</text:p></text:list-item>"
    read = read_document(
        FileType.odt,
        _odt(
            f"<text:list>{item.format('a')}</text:list>"
            f"<text:list>{item.format('b')}</text:list>"
            f'<text:list text:continue-numbering="true">{item.format("c")}</text:list>'
        ),
    )

    assert _body(read) == (
        "<ul><li><p>a</p></li></ul><ul><li><p>b</p></li><li><p>c</p></li></ul>"
    )


def test_a_table_in_an_opendocument_list_item_keeps_its_text() -> None:
    read = read_document(
        FileType.odt,
        _odt(
            "<text:list><text:list-item><text:p>Stalls</text:p>"
            "<table:table><table:table-row><table:table-cell><text:p>Cakes"
            "</text:p></table:table-cell></table:table-row></table:table>"
            "</text:list-item></text:list>"
        ),
    )

    assert _body(read) == (
        "<ul><li><p>Stalls</p><table><tr><td><p>Cakes</p></td></tr></table></li></ul>"
    )


def _repeated_table(rows: int, cells: int) -> str:
    return (
        f'<table:table><table:table-row table:number-rows-repeated="{rows}">'
        f'<table:table-cell table:number-columns-repeated="{cells}"><text:p>x'
        "</text:p></table:table-cell>"
        '<table:table-cell table:number-columns-repeated="1000"/>'
        "</table:table-row>"
        '<table:table-row table:number-rows-repeated="1000"><table:table-cell/>'
        "</table:table-row></table:table>"
    )


def test_repeated_opendocument_rows_and_cells_are_written_out() -> None:
    read = read_document(FileType.odt, _odt(_repeated_table(2, 2)))

    row = "<tr><td><p>x</p></td><td><p>x</p></td><td /></tr>"
    assert _body(read) == f"<table>{row}{row}<tr><td /></tr></table>"


@pytest.mark.parametrize(
    ("cell", "times"),
    [
        ("<table:table-cell/>", 999_999),
        # 30 cells' overhead fits; 30 copies of their text do not.
        (f"<table:table-cell><text:p>{'x' * 5000}</text:p></table:table-cell>", 30),
    ],
    ids=["many-cells", "long-text"],
)
def test_repeats_past_what_a_document_holds_are_refused(
    monkeypatch, cell, times
) -> None:
    monkeypatch.setattr(document_import, "MAX_TABLE_BYTES", 100_000)
    repeated = cell.replace(
        "<table:table-cell",
        f'<table:table-cell table:number-columns-repeated="{times}"',
        1,
    )

    with pytest.raises(FileContentError) as refused:
        read_document(
            FileType.odt,
            _odt(
                f"<table:table><table:table-row>{repeated}<table:table-cell>"
                "<text:p>end</text:p></table:table-cell></table:table-row>"
                "</table:table>"
            ),
        )

    assert refused.value.code == FileMessages.TOO_LARGE_TO_CONVERT


def test_an_opendocument_picture_keeps_its_size_and_description() -> None:
    read = read_document(
        FileType.odt,
        _odt(
            '<text:p><draw:frame svg:width="2.54cm" svg:height="1.27cm">'
            '<draw:image xlink:href="Pictures/a.png"/>'
            "<svg:desc>A red box</svg:desc></draw:frame></text:p>",
            pictures={"Pictures/a.png": _png()},
        ),
    )

    assert _body(read) == (
        '<p><ac:image ac:alt="A red box" ac:width="96" ac:height="48">'
        '<ri:attachment ri:filename="Pictures/a.png" /></ac:image></p>'
    )
    assert read.pictures == {"Pictures/a.png": _png()}


# ── Refusals ─────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "data",
    [
        b"not a zip at all",
        _docx("<w:p><w:r><w:t>unclosed"),
        _docx("").replace(
            b"<w:document",
            b'<!DOCTYPE d [<!ENTITY a "aaaa">]><w:document',
        ),
    ],
    ids=["not-a-zip", "broken-xml", "entities"],
)
def test_a_file_that_is_not_a_document_this_reads_is_refused(data: bytes) -> None:
    with pytest.raises(FileContentError) as refused:
        read_document(FileType.docx, data)

    assert refused.value.code == FileMessages.INVALID_FILE


def test_a_package_declaring_more_than_the_limit_is_refused(monkeypatch) -> None:
    monkeypatch.setattr(document_import, "MAX_DOCUMENT_PACKAGE_BYTES", 16)

    with pytest.raises(FileContentError) as refused:
        read_document(FileType.docx, _docx("<w:p/>"))

    assert refused.value.code == FileMessages.TOO_LARGE_TO_CONVERT
    assert refused.value.status_code == 413
