"""Convert to edit, through the real routes."""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from docx.shared import Inches
from httpx import AsyncClient
from PIL import Image
from sqlalchemy.exc import IntegrityError
from sqlmodel.ext.asyncio.session import AsyncSession
from sqlmodel import select

from app.models.platform.guild import CommunityRole
from app.models.tenant.file import FileVersion
from app.models.tenant.upload import Upload
from app.services.tenant import attachments, file_conversion
from app.testing import route_session_to_guild


async def _upload(client: AsyncClient, actor, filename: str, body: bytes) -> dict:
    uploaded = await client.post(
        actor.g("/files/upload"),
        headers=actor.headers,
        data={"name": "Budget", "initiative_id": str(actor.initiative.id)},
        files={"file": (filename, body, "application/octet-stream")},
    )
    assert uploaded.status_code == 201, uploaded.text
    return uploaded.json()


async def _convert(client: AsyncClient, actor, file_id: int):
    return await client.post(
        actor.g(f"/files/{file_id}/convert"), headers=actor.headers
    )


async def test_a_csv_becomes_a_spreadsheet_and_lets_its_upload_go(
    client: AsyncClient, acting_user, session
):
    owner = await acting_user(guild_role=CommunityRole.admin, initiative=True)
    uploaded = await _upload(client, owner, "budget.csv", b"Item,Cost\nTea,3\n")
    assert uploaded["converts_to"] == "sheet"

    converted = await _convert(client, owner, uploaded["id"])

    assert converted.status_code == 200, converted.text
    body = converted.json()
    assert (body["id"], body["name"]) == (uploaded["id"], "Budget")
    assert (body["file_type"], body["kind"], body["converts_to"]) == (
        "sheet",
        "spreadsheet",
        None,
    )
    (sheet,) = body["content"]["sheets"]
    assert sheet["cells"]["1:1"] == 3
    assert body["file_url"] is None
    await route_session_to_guild(session, owner.guild.id)
    assert (await session.exec(select(FileVersion))).first() is None
    assert (await session.exec(select(Upload))).first() is None
    assert await attachments.get_guild_storage_usage(owner.guild.id) == 0


@pytest.mark.parametrize(
    ("filename", "body", "first"),
    [
        ("notes.md", b"# Minutes\n\nTea was **late**.\n", "heading"),
        ("notes.txt", b"# not a heading\nsecond line\n", "paragraph"),
        ("notes.txt", "caf\u00e9 au lait\n".encode("cp1252"), "paragraph"),
    ],
    ids=["markdown", "plain-text", "windows-1252"],
)
async def test_text_becomes_a_document(
    client: AsyncClient, acting_user, filename, body, first
):
    owner = await acting_user(guild_role=CommunityRole.admin, initiative=True)
    uploaded = await _upload(client, owner, filename, body)

    converted = await _convert(client, owner, uploaded["id"])

    assert converted.status_code == 200, converted.text
    document = converted.json()
    assert document["file_type"] == "lexical"
    (top, *_) = document["content"]["root"]["children"]
    assert top["type"] == first
    assert "\ufffd" not in str(top)


def _png() -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (40, 20), "red").save(out, "PNG")
    return out.getvalue()


_PNG = _png()


def _word_with_a_picture() -> bytes:
    word = Document()
    word.add_heading("Minutes", 1)
    word.add_picture(io.BytesIO(_PNG), width=Inches(1))
    saved = io.BytesIO()
    word.save(saved)
    return saved.getvalue()


async def test_a_word_file_becomes_a_document_and_its_pictures_uploads(
    client: AsyncClient, acting_user, session
):
    owner = await acting_user(guild_role=CommunityRole.admin, initiative=True)
    uploaded = await _upload(client, owner, "minutes.docx", _word_with_a_picture())
    assert uploaded["converts_to"] == "lexical"

    converted = await _convert(client, owner, uploaded["id"])

    assert converted.status_code == 200, converted.text
    heading, shown = converted.json()["content"]["root"]["children"]
    assert heading["tag"] == "h1"
    (image,) = shown["children"]
    await route_session_to_guild(session, owner.guild.id)
    (stored,) = (await session.exec(select(Upload))).all()
    assert image["src"] == f"/uploads/{owner.guild.id}/{stored.filename}"
    assert (stored.initiative_id, stored.content_type) == (
        owner.initiative.id,
        "image/png",
    )
    assert await attachments.get_guild_storage_usage(owner.guild.id) == len(_PNG)


async def test_pictures_a_failed_conversion_stored_are_taken_back(
    client: AsyncClient, acting_user, monkeypatch
):
    owner = await acting_user(guild_role=CommunityRole.admin, initiative=True)
    uploaded = await _upload(client, owner, "minutes.docx", _word_with_a_picture())
    removed: list[str] = []
    monkeypatch.setattr(
        attachments, "delete_blobs", lambda _guild, names: removed.extend(names)
    )

    converted: list[bool] = []
    convert, commit = file_conversion.convert, AsyncSession.commit

    async def converting(*args, **kwargs):
        converted.append(True)
        return await convert(*args, **kwargs)

    async def refuse_once_converted(self):
        if converted:
            raise IntegrityError("COMMIT", None, Exception("refused"))
        await commit(self)

    monkeypatch.setattr(file_conversion, "convert", converting)
    monkeypatch.setattr(AsyncSession, "commit", refuse_once_converted)
    with pytest.raises(IntegrityError):
        await _convert(client, owner, uploaded["id"])
    monkeypatch.undo()

    (name,) = removed
    assert name.endswith(".png")


async def test_a_word_file_that_does_not_read_stays_as_it_was(
    client: AsyncClient, acting_user
):
    owner = await acting_user(guild_role=CommunityRole.admin, initiative=True)
    broken = io.BytesIO()
    with zipfile.ZipFile(broken, "w") as package:
        package.writestr("word/document.xml", "<w:document")
    uploaded = await _upload(client, owner, "minutes.docx", broken.getvalue())

    refused = await _convert(client, owner, uploaded["id"])

    assert refused.status_code == 400
    assert refused.json()["detail"] == "FILE_INVALID_FILE"
    kept = await client.get(owner.g(f"/files/{uploaded['id']}"), headers=owner.headers)
    assert kept.json()["file_type"] == "docx"


async def test_a_file_that_converts_to_nothing_is_refused(
    client: AsyncClient, acting_user
):
    owner = await acting_user(guild_role=CommunityRole.admin, initiative=True)
    uploaded = await _upload(client, owner, "report.pdf", b"%PDF-1.4 a report")
    assert uploaded["converts_to"] is None

    refused = await _convert(client, owner, uploaded["id"])

    assert refused.status_code == 400
    assert refused.json()["detail"] == "FILE_NOT_CONVERTIBLE"


async def test_a_file_too_large_to_convert_stays_as_it_was(
    client: AsyncClient, acting_user, monkeypatch
):
    owner = await acting_user(guild_role=CommunityRole.admin, initiative=True)
    uploaded = await _upload(client, owner, "budget.csv", b"Item,Cost\nTea,3\n")
    monkeypatch.setattr(file_conversion, "MAX_CONVERT_BYTES", 4)

    refused = await _convert(client, owner, uploaded["id"])

    assert refused.status_code == 413
    assert refused.json()["detail"] == "FILE_TOO_LARGE_TO_CONVERT"
    kept = await client.get(owner.g(f"/files/{uploaded['id']}"), headers=owner.headers)
    assert kept.json()["file_type"] == "csv"


async def test_only_somebody_who_may_edit_the_file_converts_it(
    client: AsyncClient, acting_user
):
    owner = await acting_user(guild_role=CommunityRole.admin, initiative=True)
    reader = await acting_user(
        guild_role=CommunityRole.member,
        guild=owner.guild,
        initiative=owner.initiative,
        initiative_role="member",
    )
    uploaded = await _upload(client, owner, "budget.csv", b"Item,Cost\nTea,3\n")

    refused = await _convert(client, reader, uploaded["id"])

    assert refused.status_code in (403, 404)
